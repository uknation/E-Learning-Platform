"""
SkillRise AI — ATS Resume Matcher Blueprint
Provides AI-driven resume cross-matching against Job Descriptions,
file parsing for PDF/DOCX/TXT, score breakdowns, and skill gap remediation.
"""

import io
import json
from flask import Blueprint, render_template, request, jsonify
from database.models import db, ATSScan, User
from services.gemini_service import analyze_ats_resume, improve_single_bullet
from utils import get_current_user

ats_bp = Blueprint('ats', __name__)

SAMPLE_JD_PRESETS = [
    {
        "id": "fullstack",
        "title": "Full Stack Developer",
        "icon": "layers",
        "jd": """We are looking for a Full Stack Developer proficient in Python/Flask or Node.js, and modern frontend frameworks like React or Vue.js. 
Requirements:
- Strong knowledge of Data Structures, Algorithms, and Object-Oriented Programming.
- Hands-on experience designing and consuming RESTful APIs.
- Database expertise with SQL (PostgreSQL/MySQL) and NoSQL (MongoDB).
- Familiarity with Git version control, Docker containers, and CI/CD pipelines.
- Experience writing unit tests, debugging, and optimizing web application performance.
- Excellent communication, teamwork, and problem-solving abilities."""
    },
    {
        "id": "frontend",
        "title": "Frontend Engineer",
        "icon": "layout",
        "jd": """Seeking a passionate Frontend Engineer specializing in modern JavaScript/TypeScript and React.js.
Key Responsibilities:
- Build responsive, accessible, high-performance UI components from Figma designs.
- Manage client state effectively with Redux Toolkit or React Context.
- Optimize web vitals, page load speeds, and cross-browser compatibility.
- Integrate backend REST & WebSocket APIs.
- Write unit & integration tests using Jest and React Testing Library.
- Collaborate with product managers and backend engineers in an Agile/Scrum environment."""
    },
    {
        "id": "backend_python",
        "title": "Backend Engineer",
        "icon": "terminal",
        "jd": """Looking for a Backend Engineer to build scalable distributed microservices and robust API pipelines.
Requirements:
- Strong Python/Backend mastery (FastAPI, Flask, Django, Node.js) and asynchronous programming.
- Experience with REST API integrations and backend services architecture.
- Relational database schema design and query optimization (SQLAlchemy, PostgreSQL, MySQL).
- Experience with cloud infrastructure (AWS/GCP), Redis caching, and Docker.
- Understanding of microservices architecture, message queues, and API security."""
    },
    {
        "id": "data_analyst",
        "title": "Data Analyst",
        "icon": "bar-chart-2",
        "jd": """Join our analytics team to extract actionable insights from complex datasets.
Requirements:
- Advanced SQL querying skills (window functions, CTEs, complex joins).
- Proficiency in Python for data manipulation (Pandas, NumPy) and visualization (Matplotlib, Seaborn).
- Experience building interactive dashboards in PowerBI or Tableau.
- Strong mathematical, statistical, and quantitative aptitude.
- Ability to communicate data findings clearly to non-technical business stakeholders."""
    }
]


def extract_text_from_file(file_storage) -> str:
    """Extract text from uploaded PDF, DOCX, or TXT file."""
    filename = (file_storage.filename or '').lower()
    
    if filename.endswith('.pdf'):
        try:
            import pypdf
            reader = pypdf.PdfReader(file_storage.stream)
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            return "\n\n".join(text_parts).strip()
        except Exception as e:
            print(f"Error parsing PDF with pypdf: {e}")
            return ""
            
    elif filename.endswith('.docx'):
        try:
            import docx
            doc = docx.Document(file_storage.stream)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        full_text.append(row_text)
            return "\n".join(full_text).strip()
        except Exception as e:
            print(f"Error parsing DOCX: {e}")
            return ""
            
    elif filename.endswith('.txt'):
        try:
            content = file_storage.read()
            try:
                return content.decode('utf-8').strip()
            except UnicodeDecodeError:
                return content.decode('latin-1', errors='ignore').strip()
        except Exception as e:
            print(f"Error reading TXT file: {e}")
            return ""
            
    return ""


@ats_bp.route('/ats-matcher')
def ats_page():
    user = get_current_user()
    return render_template('ats_matcher.html', presets=SAMPLE_JD_PRESETS, user=user)


@ats_bp.route('/api/ats/analyze', methods=['POST'])
def analyze_resume():
    user = get_current_user()
    
    resume_text = ""
    job_description = ""
    target_role = ""
    
    # Handle Multipart Form (File Upload)
    if request.files and 'resume_file' in request.files:
        uploaded_file = request.files['resume_file']
        if uploaded_file and uploaded_file.filename != '':
            resume_text = extract_text_from_file(uploaded_file)
            
        job_description = request.form.get('job_description', '').strip()
        target_role = request.form.get('target_role', '').strip()
        
        # If resume text was also pasted as fallback
        if not resume_text and request.form.get('resume_text'):
            resume_text = request.form.get('resume_text', '').strip()
            
    # Handle JSON payload
    elif request.is_json:
        data = request.get_json() or {}
        resume_text = data.get('resume_text', '').strip()
        job_description = data.get('job_description', '').strip()
        target_role = data.get('target_role', '').strip()
        
    else:
        resume_text = request.form.get('resume_text', '').strip()
        job_description = request.form.get('job_description', '').strip()
        target_role = request.form.get('target_role', '').strip()

    if not resume_text:
        return jsonify({
            'error': 'Please provide resume text or upload a valid PDF, DOCX, or TXT file.'
        }), 400

    if not job_description:
        return jsonify({
            'error': 'Please provide a target Job Description to cross-match against.'
        }), 400

    # Call AI ATS Analysis
    analysis = analyze_ats_resume(resume_text, job_description, target_role)
    
    # Award user XP
    try:
        if user:
            user.add_xp(25)  # 25 XP for conducting an ATS resume audit
            
            # Save record in database
            scan_record = ATSScan(
                user_id=user.id,
                target_role=analysis.get('target_role', target_role or 'General'),
                score=int(analysis.get('overall_score', 0)),
                matched_skills=json.dumps(analysis.get('matched_skills', [])),
                missing_skills=json.dumps(analysis.get('missing_hard_skills', []) + analysis.get('missing_soft_skills', [])),
                analysis_json=json.dumps(analysis)
            )
            db.session.add(scan_record)
            db.session.commit()
    except Exception as e:
        print(f"Error saving ATS scan or awarding XP: {e}")

    return jsonify({
        'status': 'success',
        'analysis': analysis,
        'xp_awarded': 25
    })


@ats_bp.route('/api/ats/improve-bullet', methods=['POST'])
def improve_bullet():
    data = request.get_json() or {}
    bullet_text = data.get('bullet_text', '').strip()
    target_role = data.get('target_role', '').strip()
    
    if not bullet_text:
        return jsonify({'error': 'Bullet text cannot be empty'}), 400
        
    result = improve_single_bullet(bullet_text, target_role)
    return jsonify({
        'status': 'success',
        'result': result
    })


@ats_bp.route('/api/ats/history')
def get_history():
    user = get_current_user()
    try:
        scans = ATSScan.query.filter_by(user_id=user.id).order_by(ATSScan.created_at.desc()).limit(10).all()
        return jsonify({
            'status': 'success',
            'scans': [s.to_dict() for s in scans]
        })
    except Exception as e:
        print(f"Error loading ATS scan history: {e}")
        return jsonify({'status': 'success', 'scans': []})
